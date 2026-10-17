import io
import socket
import unittest
from datetime import timedelta
from unittest.mock import patch

from tests.support import Fixture
from app.models import (Submission, SubmissionRecord, WeekExport,
                        WeekConfig, AttendanceTask, AttendanceRecord, DutyAssignment, Student, RankingSummary)
from app.tasks.ranking_calculator import calculate_class_stats, get_period_range, get_previous_period_range
from openpyxl import load_workbook
from app.tasks import ranking_calculator


class BackendRegressionTest(unittest.TestCase):
    def setUp(self):
        connect = socket.socket.connect
        def no_external(sock, address):
            # Windows asyncio implements socketpair via localhost TCP.
            if isinstance(address, tuple) and address[0] in ("127.0.0.1", "::1"):
                return connect(sock, address)
            raise AssertionError("External network forbidden in backend tests")
        self.network = patch("socket.socket.connect", no_external)
        self.network.start()
        self.fx = Fixture()
        self.client = self.fx.client

    def tearDown(self):
        self.fx.close()
        self.network.stop()

    def submit(self, tasks=None, week=1, **extra):
        return self.client.post("/api/submissions/", headers=self.fx.headers(),
            json={"week_number": week, "task_ids": tasks or ["task-1"], **extra})

    def seed_submission(self, sid=1, semester=2, week=1, status="pending", record_ids=(1, 2)):
        self.fx.db.add(Submission(id=sid, semester_id=semester, user_id=1, week_number=week,
                                 status=status, class_names="Test01"))
        self.fx.db.commit()
        self.fx.db.add_all([SubmissionRecord(submission_id=sid, record_id=rid) for rid in record_ids])
        self.fx.db.commit()

    def test_authentication_required_on_every_task_and_record_path(self):
        requests = [("get", "/api/tasks", None), ("get", "/api/tasks/task-1", None),
            ("post", "/api/tasks", {"id": "new", "type": "name_check", "class_ids": [1]}),
            ("put", "/api/tasks/task-1", {"status": "abandoned"}),
            ("get", "/api/tasks/task-1/records", None),
            ("post", "/api/tasks/task-1/records", [{"student_id": 1, "class_id": 1, "status": "present"}]),
            ("put", "/api/records/1", {"status": "present"}),
            ("post", "/api/records/batch-update", [{"task_id": "task-1", "student_id": 1, "status": "present"}]),
            ("delete", "/api/records/by-task-student?task_id=task-1&student_id=1", None)]
        for method, path, body in requests:
            with self.subTest(path=path, method=method):
                response = self.client.request(method, path, json=body)
                self.assertEqual(response.status_code, 401, response.text)
        self.assertEqual(self.client.put("/api/tasks/task-1", headers={"Authorization": "Bearer invalid"}, json={"status": "abandoned"}).status_code, 401)

    def test_other_account_and_unowned_legacy_task_are_denied(self):
        for path in ("/api/tasks/task-1", "/api/tasks/task-1/records", "/api/tasks/legacy-task"):
            self.assertEqual(self.client.get(path, headers=self.fx.headers(3)).status_code, 403)
        self.assertEqual(self.client.put("/api/records/1", headers=self.fx.headers(3), json={"status": "present"}).status_code, 403)

    def test_post_and_put_share_submission_lock_and_idempotent_retry(self):
        self.seed_submission()
        payload = {"id": "task-1", "type": "name_check", "class_ids": [1], "status": "abandoned"}
        self.assertEqual(self.client.post("/api/tasks", headers=self.fx.headers(), json=payload).status_code, 403)
        self.assertEqual(self.client.put("/api/tasks/task-1", headers=self.fx.headers(), json={"status": "abandoned"}).status_code, 403)
        payload["status"] = "completed"
        self.assertEqual(self.client.post("/api/tasks", headers=self.fx.headers(), json=payload).status_code, 201)
        self.assertEqual(self.client.delete("/api/records/by-task-student?task_id=task-1&student_id=1", headers=self.fx.headers()).status_code, 403)
        self.assertEqual(self.client.post("/api/tasks/task-1/records", headers=self.fx.headers(), json=[{"student_id": 1, "class_id": 1, "status": "present"}]).status_code, 403)

    def test_record_delete_is_idempotent_and_has_owner_check(self):
        path = "/api/records/by-task-student?task_id=task-1&student_id=1"
        self.assertEqual(self.client.delete(path, headers=self.fx.headers(3)).status_code, 403)
        self.assertTrue(self.client.delete(path, headers=self.fx.headers()).json()["deleted"])
        self.assertFalse(self.client.delete(path, headers=self.fx.headers()).json()["deleted"])

    def test_approved_completed_task_allows_only_final_workflow_metadata(self):
        self.seed_submission(status="approved")
        for method in ("put", "post"):
            with self.subTest(method=method):
                task = self.fx.db.get(AttendanceTask, "task-1")
                task.phase = "confirming"
                self.fx.db.commit()
                body = {"phase": "finished"}
                path = "/api/tasks/task-1"
                if method == "post":
                    path = "/api/tasks"
                    body.update(id="task-1", type="name_check", class_ids=[1], status="completed")
                denied = self.client.request(method, path, headers=self.fx.headers(3), json=body)
                self.assertEqual(denied.status_code, 403)
                response = self.client.request(method, path, headers=self.fx.headers(), json=body)
                self.assertEqual(response.status_code, 200 if method == "put" else 201, response.text)
                self.assertEqual(response.json()["phase"], "finished")
                self.assertEqual(response.json()["status"], "completed")
                self.fx.db.expire_all()
                self.assertEqual(self.fx.db.get(AttendanceRecord, 1).status, "absent")
                self.assertEqual(self.fx.db.get(Submission, 1).status, "approved")

    def test_completion_metadata_cannot_bypass_approved_task_mutation_protection(self):
        self.seed_submission(status="approved")
        task = self.fx.db.get(AttendanceTask, "task-1")
        task.phase = "confirming"
        self.fx.db.commit()
        for patch_body in ({"status": "abandoned"}, {"status": "in_progress"},
                           {"current_class_index": 1}, {"current_student_index": 1},
                           {"phase": "executing"}, {"phase": "finished", "status": "abandoned"},
                           {"phase": "finished", "current_student_index": 1}):
            for method in ("put", "post"):
                with self.subTest(method=method, body=patch_body):
                    body = dict(patch_body)
                    path = "/api/tasks/task-1"
                    if method == "post":
                        path = "/api/tasks"
                        body.update(id="task-1", type="name_check", class_ids=[1])
                    response = self.client.request(method, path, headers=self.fx.headers(), json=body)
                    self.assertEqual(response.status_code, 403, response.text)
        self.fx.db.expire_all()
        task = self.fx.db.get(AttendanceTask, "task-1")
        self.assertEqual(task.status, "completed")
        self.assertEqual(task.phase, "confirming")
        task.status = "in_progress"
        self.fx.db.commit()
        self.assertEqual(self.client.put("/api/tasks/task-1", headers=self.fx.headers(),
            json={"phase": "finished"}).status_code, 403)

    def test_explicit_null_clears_remark_while_omission_keeps_it(self):
        self.assertEqual(self.client.put("/api/records/1", headers=self.fx.headers(), json={"status": "present"}).json()["remark"], "old note")
        self.assertIsNone(self.client.put("/api/records/1", headers=self.fx.headers(), json={"status": "present", "remark": None}).json()["remark"])
        self.client.put("/api/records/1", headers=self.fx.headers(), json={"status": "other", "remark": "new note"})
        result = self.client.post("/api/tasks/task-1/records", headers=self.fx.headers(), json=[{"student_id": 1, "class_id": 1, "status": "present", "remark": None}])
        self.assertIsNone(result.json()[0]["remark"])
        self.client.put("/api/records/1", headers=self.fx.headers(), json={"status": "other", "remark": "again"})
        self.client.post("/api/records/batch-update", headers=self.fx.headers(), json=[{"task_id": "task-1", "student_id": 1, "status": "present", "remark": None}])
        self.assertIsNone(self.client.get("/api/tasks/task-1/records", headers=self.fx.headers()).json()[0]["remark"])

    def test_submission_detail_success_and_semester_scope(self):
        self.seed_submission(semester=1, status="approved")
        self.fx.db.add(WeekExport(semester_id=1, week_number=1, exported_by=2)); self.fx.db.commit()
        response = self.client.get("/api/submissions/1", headers=self.fx.headers())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["class_names"], "Test01")
        summary = self.client.get("/api/submissions/week-summary/1", headers=self.fx.headers()).json()
        self.assertEqual(summary["total_submissions"], 0)
        self.assertFalse(summary["is_published"])
        self.assertEqual(self.client.get("/api/submissions/week-summary-detail/1", headers=self.fx.headers()).status_code, 403)
        old = self.client.get("/api/submissions/week-summary/1?semester_id=1", headers=self.fx.headers()).json()
        self.assertEqual(old["approved_count"], 1)
        self.assertTrue(old["is_published"])

    def test_new_semester_keeps_old_identity_and_validates_monday(self):
        old_start = self.fx.db.get(WeekConfig, 2).start_date
        next_start = self.fx.start + timedelta(weeks=24)
        response = self.client.put("/api/week/config", headers=self.fx.headers(2), json={"start_date": next_start.isoformat(), "semester_name": "next"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotEqual(response.json()["semester_id"], 2)
        self.fx.db.expire_all()
        self.assertEqual(self.fx.db.get(WeekConfig, 2).start_date, old_start)
        self.assertFalse(self.fx.db.get(WeekConfig, 2).is_active)
        self.assertEqual(self.client.put("/api/week/config", headers=self.fx.headers(2), json={"start_date": (next_start + timedelta(days=1)).isoformat()}).status_code, 422)

    def test_week_validation_allows_past_and_rejects_negative_future_wrong_term(self):
        self.assertEqual(self.submit(week=-99).status_code, 422)
        future = (self.fx.now.date() - self.fx.start).days // 7 + 2
        self.assertEqual(self.submit(week=future).status_code, 400)
        self.assertEqual(self.submit(week=1, semester_id=999).status_code, 404)
        self.assertEqual(self.submit(week=1, semester_id=1).status_code, 400)
        accepted = self.submit(week=1)
        self.assertEqual(accepted.status_code, 200, accepted.text)
        self.assertEqual(accepted.json()["semester_id"], 2)

    def test_leave_and_other_export_rows_and_formula_columns(self):
        self.fx.db.get(AttendanceRecord, 1).status = "leave"
        self.fx.db.get(AttendanceRecord, 2).status = "other"; self.fx.db.commit()
        created = self.submit(); self.assertEqual(created.status_code, 200, created.text)
        approved = self.client.put(f"/api/submissions/{created.json()['id']}/approve", headers=self.fx.headers(2))
        self.assertEqual(approved.status_code, 200, approved.text)
        response = self.client.get("/api/submissions/export/1", headers=self.fx.headers(2))
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else "")
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet.max_column, 10)
        self.assertEqual([sheet.cell(1, i).value for i in (7, 8)], ["请假", "其他"])
        self.assertEqual(sum(sheet.cell(row, 7).value for row in (2, 3)), 1)
        self.assertEqual(sum(sheet.cell(row, 8).value for row in (2, 3)), 1)
        self.assertEqual(sheet.cell(2, 9).value, "=ROUNDDOWN(E2/2+F2,0)")

    def test_ranking_counts_tasks_and_preserves_approved_roster_size(self):
        created = self.submit(tasks=["task-1", "task-2"])
        self.assertEqual(created.status_code, 200, created.text)
        self.client.put(f"/api/submissions/{created.json()['id']}/approve", headers=self.fx.headers(2))
        self.fx.db.add(Student(id=3, name="Later Student", student_no="003", class_id=1)); self.fx.db.commit()
        stats = calculate_class_stats(self.fx.db, self.fx.now - timedelta(days=1), self.fx.now + timedelta(days=1), "rate", 2)
        self.assertEqual(stats[1]["total_expected"], 4)
        self.assertEqual(stats[1]["rate"], .5)
        self.assertEqual(calculate_class_stats(self.fx.db, self.fx.now - timedelta(days=1), self.fx.now + timedelta(days=1), "rate", 1), {})

    def test_ranking_windows_are_half_open_and_equal_length(self):
        for period, length in (("7d", 7), ("30d", 30)):
            start, end = get_period_range(period, self.fx.now.date())
            prev_start, prev_end = get_previous_period_range(period, self.fx.now.date())
            self.assertEqual(end - start, timedelta(days=length))
            self.assertEqual(prev_end - prev_start, timedelta(days=length))
            self.assertEqual(prev_end, start)

    def test_failed_ranking_rebuild_keeps_last_good_cache(self):
        self.fx.db.add(RankingSummary(semester_id=2, period_type="7d", rank_type="rate", avg_value=.25, total_classes=1))
        self.fx.db.commit()
        with patch.object(ranking_calculator, "SessionLocal", self.fx.sessions), patch.object(
            ranking_calculator, "calculate_class_stats", side_effect=RuntimeError("synthetic failure")):
            with self.assertLogs(ranking_calculator.logger, level="ERROR"):
                with self.assertRaisesRegex(RuntimeError, "synthetic failure"):
                    ranking_calculator.calculate_rankings()
        self.fx.db.expire_all()
        self.assertEqual(float(self.fx.db.query(RankingSummary).one().avg_value), .25)

    def test_ranking_rebuild_does_not_remove_other_semester_cache(self):
        self.fx.db.add(RankingSummary(semester_id=1, period_type="7d", rank_type="rate", avg_value=.75, total_classes=1))
        self.fx.db.commit()
        with patch.object(ranking_calculator, "SessionLocal", self.fx.sessions):
            ranking_calculator.calculate_rankings()
        self.fx.db.expire_all()
        self.assertEqual(float(self.fx.db.query(RankingSummary).filter_by(semester_id=1).one().avg_value), .75)
        self.assertEqual(self.fx.db.query(RankingSummary).filter_by(semester_id=2).count(), 9)

    def test_every_week_list_path_defaults_to_active_semester(self):
        self.seed_submission(semester=1, status="pending")
        for path in ("/api/submissions/", "/api/submissions/pending", "/api/submissions/reviewed"):
            result = self.client.get(path, headers=self.fx.headers(2)).json()
            self.assertEqual(result, [], path)
        result = self.client.get("/api/submissions/admin-search", headers=self.fx.headers(2)).json()
        self.assertEqual(result["total"], 0)
        self.assertEqual(self.client.get("/api/submissions/submitted-task-ids", headers=self.fx.headers()).json(), {"task_ids": []})

    def test_cancelled_and_old_semester_submissions_do_not_satisfy_duty(self):
        self.fx.db.add(DutyAssignment(user_id=1, assigned_by=2, is_active=True)); self.fx.db.commit()
        self.seed_submission(status="cancelled", record_ids=())
        self.seed_submission(sid=2, semester=1, status="approved", record_ids=())
        response = self.client.get("/api/duties/week-submissions?week_number=1", headers=self.fx.headers(2))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["submitted_count"], 0)
        self.assertEqual(response.json()["not_submitted_count"], 1)
        missing = self.client.get("/api/duties/unsubmitted?week_number=1", headers=self.fx.headers(2)).json()
        self.assertEqual([item["user_id"] for item in missing], [1])


if __name__ == "__main__":
    unittest.main()
